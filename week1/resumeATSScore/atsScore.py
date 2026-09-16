import os
from pathlib import Path
from dotenv import load_dotenv
from groq import Groq

import pdfplumber
from docx import Document

from pathlib import Path

load_dotenv()
my_api_key=os.getenv("GROQ_API_KEY")

if not my_api_key:
    raise ValueError("API key missed")

client=Groq(api_key=my_api_key)
model="openai/gpt-oss-120b"
role="user"

class UnsupportedFileType(ValueError):
    """Raised when the file extension isn't .pdf or .docx."""

class EmptyExtractionError(ValueError):
    """Raised when extraction succeeds but yields no usable text.
    Usually means the PDF is a scanned image (no embedded text layer),
    which this module does not OCR.
    """


def _extract_pdf(path: Path) -> str:
    chunks: list[str] = []
    with pdfplumber.open(path) as pdf:
        for page in pdf.pages:
            page_text = page.extract_text() or ""
            if page_text.strip():
                chunks.append(page_text)

            # Tables often hold skills/dates that extract_text() misses
            # or garbles (e.g. skill matrices, education tables).
            for table in page.extract_tables():
                for row in table:
                    cells = [c.strip() for c in row if c and c.strip()]
                    if cells:
                        chunks.append(" | ".join(cells))

    return "\n".join(chunks)


def _extract_docx(path: Path) -> str:
    doc = Document(path)
    chunks: list[str] = []

    for para in doc.paragraphs:
        if para.text.strip():
            chunks.append(para.text)

    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                chunks.append(" | ".join(cells))

    return "\n".join(chunks)


def extract_text(path: str | Path) -> str:
    """Extract plain text from a resume file (.pdf or .docx).

    Raises:
        FileNotFoundError: if `path` doesn't exist.
        UnsupportedFileType: if the extension isn't .pdf or .docx.
        EmptyExtractionError: if extraction ran but found no text
            (common for scanned/image-only PDFs — needs OCR, out of
            scope for this module).
    """
    path = Path(path)
    if not path.exists():
        raise FileNotFoundError(f"No such file: {path}")

    suffix = path.suffix.lower()
    if suffix == ".pdf":
        text = _extract_pdf(path)
    elif suffix == ".docx":
        text = _extract_docx(path)
    else:
        raise UnsupportedFileType(
            f"Unsupported file type '{suffix}'. Expected .pdf or .docx."
        )

    if not text.strip():
        raise EmptyExtractionError(
            f"No extractable text found in {path}. "
            "If this is a scanned/image PDF, it needs OCR first."
        )

    return text